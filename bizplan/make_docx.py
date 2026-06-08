from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import re

doc = Document()

# ── Page margins ──────────────────────────────────────────────────────────────
for section in doc.sections:
    section.top_margin    = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin   = Cm(2.5)
    section.right_margin  = Cm(2)

# ── Styles ────────────────────────────────────────────────────────────────────
styles = doc.styles

def set_style(name, font_name='Calibri', size=11, bold=False, color=None, space_before=0, space_after=6):
    try:
        s = styles[name]
    except KeyError:
        return
    f = s.font
    f.name = font_name
    f.size = Pt(size)
    f.bold = bold
    if color:
        f.color.rgb = RGBColor(*color)
    pf = s.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after  = Pt(space_after)

set_style('Normal',    size=11, space_after=4)
set_style('Heading 1', size=16, bold=True,  color=(31, 73, 125),  space_before=14, space_after=6)
set_style('Heading 2', size=13, bold=True,  color=(47, 84, 150),  space_before=10, space_after=4)
set_style('Heading 3', size=11, bold=True,  color=(68, 114, 196), space_before=8,  space_after=3)

# ── Helper: shade table header row ───────────────────────────────────────────
def shade_row(row, fill='1F497D'):
    for cell in row.cells:
        tc   = cell._tc
        tcPr = tc.get_or_add_tcPr()
        shd  = OxmlElement('w:shd')
        shd.set(qn('w:val'),   'clear')
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:fill'),  fill)
        tcPr.append(shd)
        for para in cell.paragraphs:
            for run in para.runs:
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                run.font.bold = True

def set_cell_bg(cell, fill='EBF3FB'):
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd  = OxmlElement('w:shd')
    shd.set(qn('w:val'),   'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'),  fill)
    tcPr.append(shd)

def add_horizontal_rule(doc):
    p    = doc.add_paragraph()
    pPr  = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bot  = OxmlElement('w:bottom')
    bot.set(qn('w:val'),   'single')
    bot.set(qn('w:sz'),    '6')
    bot.set(qn('w:space'), '1')
    bot.set(qn('w:color'), '1F497D')
    pBdr.append(bot)
    pPr.append(pBdr)

# ── Parse inline bold/normal text ────────────────────────────────────────────
def add_run_with_inline(para, text):
    parts = re.split(r'\*\*(.+?)\*\*', text)
    for i, part in enumerate(parts):
        run = para.add_run(part)
        if i % 2 == 1:
            run.bold = True

# ── Parse markdown table ──────────────────────────────────────────────────────
def parse_md_table(lines):
    rows = []
    for line in lines:
        if re.match(r'^\s*\|?[-:| ]+\|?\s*$', line):
            continue
        cells = [c.strip() for c in re.split(r'\|', line) if c.strip() != '']
        if cells:
            rows.append(cells)
    return rows

def add_table(doc, rows):
    if not rows:
        return
    cols = max(len(r) for r in rows)
    t    = doc.add_table(rows=len(rows), cols=cols)
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    for ri, row in enumerate(rows):
        for ci, cell_text in enumerate(row):
            cell = t.cell(ri, ci)
            cell.text = ''
            para = cell.paragraphs[0]
            para.paragraph_format.space_before = Pt(2)
            para.paragraph_format.space_after  = Pt(2)
            add_run_with_inline(para, cell_text)
            for run in para.runs:
                run.font.size = Pt(10)
        if ri == 0:
            shade_row(t.rows[0])
        elif ri % 2 == 0:
            for ci in range(cols):
                set_cell_bg(t.cell(ri, ci), 'EBF3FB')
    doc.add_paragraph()

# ── Main parse loop ────────────────────────────────────────────────────────────
with open('bizplan_sotskontrakt.md', encoding='utf-8') as f:
    raw = f.read()

lines = raw.splitlines()
i = 0
while i < len(lines):
    line = lines[i]

    # Blockquote → highlighted paragraph
    if line.startswith('> '):
        p = doc.add_paragraph(style='Normal')
        p.paragraph_format.left_indent = Cm(0.8)
        add_run_with_inline(p, line[2:])
        for run in p.runs:
            run.font.color.rgb = RGBColor(31, 73, 125)
            run.font.bold = True
        i += 1
        continue

    # Heading 1 (#)
    if line.startswith('# ') and not line.startswith('## '):
        p = doc.add_heading(line[2:], level=1)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        i += 1
        continue

    # Heading 2 (##)
    if line.startswith('## ') and not line.startswith('### '):
        p = doc.add_heading(line[3:], level=2)
        i += 1
        continue

    # Heading 3 (###)
    if line.startswith('### '):
        doc.add_heading(line[4:], level=3)
        i += 1
        continue

    # Horizontal rule
    if line.strip() in ('---', '***', '___'):
        add_horizontal_rule(doc)
        i += 1
        continue

    # Table (collect all consecutive table lines)
    if line.startswith('|') or (i + 1 < len(lines) and lines[i+1].startswith('|---')):
        table_lines = []
        while i < len(lines) and (lines[i].startswith('|') or re.match(r'^\s*\|?[-:| ]+\|?\s*$', lines[i])):
            table_lines.append(lines[i])
            i += 1
        add_table(doc, parse_md_table(table_lines))
        continue

    # Bullet list
    if line.startswith('- '):
        p = doc.add_paragraph(style='List Bullet')
        add_run_with_inline(p, line[2:])
        for run in p.runs:
            run.font.size = Pt(11)
        i += 1
        continue

    # Numbered list
    m = re.match(r'^\d+\.\s+(.*)', line)
    if m:
        p = doc.add_paragraph(style='List Number')
        add_run_with_inline(p, m.group(1))
        i += 1
        continue

    # Bold-only metadata lines (**Key:** value)
    if line.startswith('**') and ':**' in line:
        p = doc.add_paragraph(style='Normal')
        add_run_with_inline(p, line)
        i += 1
        continue

    # Empty line
    if line.strip() == '':
        i += 1
        continue

    # Italic/em line (*text*)
    if line.strip().startswith('*') and line.strip().endswith('*'):
        p = doc.add_paragraph(style='Normal')
        run = p.add_run(line.strip().strip('*'))
        run.italic = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(128, 128, 128)
        i += 1
        continue

    # Normal paragraph
    p = doc.add_paragraph(style='Normal')
    add_run_with_inline(p, line)
    i += 1

doc.save('bizplan_sotskontrakt.docx')
print('OK: bizplan_sotskontrakt.docx')
